"""Zaxira (backup) fayllarini tayyorlash.

build_files() ikkita fayl qaytaradi:
  1) .xlsx - odam o'qishi uchun chiroyli, ko'p varaqli Excel (Toshkent vaqti, UTC+5)
  2) .json - bazani to'liq tiklash uchun xom nusxa (vaqtlar UTC, bazada saqlanganidek)
"""
import io
import json
from collections import Counter
from datetime import date, datetime, timedelta, timezone

LOCAL_OFFSET = timedelta(hours=5)      # Toshkent = UTC+5
MAX_JSON_BYTES = 45 * 1024 * 1024      # Telegram botlar uchun fayl limiti 50 MB


def utc_now():
    return datetime.now(timezone.utc).replace(tzinfo=None)


def to_local(dt):
    """Bazadagi UTC vaqtni Toshkent vaqtiga o'giradi (Excel uchun oddiy datetime)."""
    if dt is None:
        return None
    if isinstance(dt, str):
        try:
            dt = datetime.fromisoformat(dt)
        except ValueError:
            return dt
    return dt + LOCAL_OFFSET if isinstance(dt, datetime) else dt


def is_active_token(tok):
    return bool(tok) and tok != "NO_TOKEN" and not str(tok).startswith("ERROR")


def owner_label(user, fallback_id=None):
    """@username -> telefon -> login (admin panelidagi bilan bir xil tartib)."""
    if not user:
        return f"O'chirilgan foydalanuvchi (ID {fallback_id})" if fallback_id is not None else "—"
    if user.get("username"):
        return "@" + user["username"]
    if user.get("phone_number"):
        return user["phone_number"]
    return user.get("login") or f"ID {user.get('id')}"


def summarize(data, now_utc=None):
    now_utc = now_utc or utc_now()
    users, accs = data["app_users"], data["hero_accounts"]
    email_count = Counter((a.get("email") or "").strip().lower() for a in accs)
    active_accs = sum(1 for a in accs if is_active_token(a.get("bearer_token")))
    return {
        "users": len(users),
        "active_users": sum(1 for u in users if u.get("trial_ends_at") and u["trial_ends_at"] > now_utc),
        "unique_accounts": len(email_count),
        "account_records": len(accs),
        "duplicate_emails": sum(1 for n in email_count.values() if n > 1),
        "active_accounts": active_accs,
        "inactive_accounts": len(accs) - active_accs,
        "archived": len(data["archived_accounts"]),
        "scans": len(data["scan_logs"]),
        "share_codes": len(data["share_codes"]),
    }


def build_xlsx(data, now_utc, now_local, stats):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.properties import PageSetupProperties

    head_fill = PatternFill("solid", fgColor="3B2F25")
    head_font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    zebra_fill = PatternFill("solid", fgColor="F7F1E8")
    warn_fill = PatternFill("solid", fgColor="FFF3CD")
    thin = Side(style="thin", color="E2D8C6")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    good_font = Font(name="Calibri", bold=True, color="2E7D32")
    bad_font = Font(name="Calibri", bold=True, color="C62828")
    GOOD, BAD = {"Aktiv", "Faol"}, {"Nofaol", "Muddati tugagan"}

    def write_table(ws, headers, rows, widths, top=1, status_col=None, warn_col=None, decorate=True):
        for c, h in enumerate(headers, 1):
            cell = ws.cell(row=top, column=c, value=h)
            cell.fill, cell.font, cell.border = head_fill, head_font, border
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[top].height = 28
        for r, row in enumerate(rows, top + 1):
            for c, val in enumerate(row, 1):
                cell = ws.cell(row=r, column=c, value=val)
                cell.border = border
                cell.alignment = Alignment(vertical="center", horizontal="left" if isinstance(val, str) else "center")
                if (r - top) % 2 == 0:
                    cell.fill = zebra_fill
                if isinstance(val, datetime):
                    cell.number_format = "DD.MM.YYYY HH:MM:SS"
                elif isinstance(val, int) and not isinstance(val, bool):
                    cell.number_format = "0"
                elif isinstance(val, float):
                    cell.number_format = "0.00"
                if status_col and c == status_col:
                    if val in GOOD:
                        cell.font = good_font
                    elif val in BAD:
                        cell.font = bad_font
                if warn_col and c == warn_col and val not in (None, "", "—"):
                    cell.fill = warn_fill
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        # Bosib chiqarish/PDF: har varaq eniga bir sahifaga sig'sin
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        if decorate:
            ws.freeze_panes = ws.cell(row=top + 1, column=1)
            ws.auto_filter.ref = f"A{top}:{get_column_letter(len(headers))}{max(len(rows) + top, top)}"

    users, accs = data["app_users"], data["hero_accounts"]
    scans, arch, shares = data["scan_logs"], data["archived_accounts"], data["share_codes"]
    by_id = {u["id"]: u for u in users}
    email_count = Counter((a.get("email") or "").strip().lower() for a in accs)
    accs_per_user = Counter(a["user_id"] for a in accs)
    scans_per_user = Counter(s["user_id"] for s in scans)

    wb = Workbook()

    # ---- 1) Umumiy ----
    ws = wb.active
    ws.title = "Umumiy"
    ws.sheet_view.showGridLines = False
    ws["A1"] = "HeroScanner — to'liq zaxira nusxa"
    ws["A1"].font = Font(name="Calibri", bold=True, size=18, color="3B2F25")
    ws["A2"] = f"Yaratilgan: {now_local:%d.%m.%Y %H:%M:%S}  (Toshkent vaqti, UTC+5)"
    ws["A2"].font = Font(name="Calibri", italic=True, size=11, color="7A6A55")
    summary = [
        ["Foydalanuvchilar (jami)", stats["users"]],
        ["Obunasi faol foydalanuvchilar", stats["active_users"]],
        ["Hero akkauntlar — noyob (takrorlarsiz)", stats["unique_accounts"]],
        ["Hero akkaunt yozuvlari (takrorlar bilan)", stats["account_records"]],
        ["Bir nechta foydalanuvchida takrorlangan emaillar", stats["duplicate_emails"]],
        ["Aktiv akkauntlar (tokeni bor)", stats["active_accounts"]],
        ["Nofaol akkauntlar", stats["inactive_accounts"]],
        ["Arxivdagi (o'chirilgan) akkauntlar", stats["archived"]],
        ["Skanerlar tarixi (yozuvlar soni)", stats["scans"]],
        ["Ulashish kodlari", stats["share_codes"]],
    ]
    write_table(ws, ["Ko'rsatkich", "Qiymat"], summary, [52, 16], top=4, decorate=False)
    guide_row = 4 + len(summary) + 2
    ws.cell(row=guide_row, column=1, value="Varaqlar:").font = Font(name="Calibri", bold=True, size=12, color="3B2F25")
    guide = [
        "Foydalanuvchilar — barcha foydalanuvchilar, obuna muddati, akkaunt va skaner soni",
        "Hero akkauntlar — har bir foydalanuvchining Hero email/paroli va holati (takrorlar belgilangan)",
        "Skaner tarixi — barcha skanerlar, natija va vaqti",
        "Arxiv — o'chirilgan akkauntlar",
        "Ulashish kodlari — yaratilgan va ishlatilgan kodlar",
    ]
    for i, line in enumerate(guide, guide_row + 1):
        ws.cell(row=i, column=1, value="• " + line).font = Font(name="Calibri", size=11, color="4A3D30")
    ws.cell(row=guide_row + len(guide) + 2, column=1,
            value="⚠️ Faylda parollar bor — hech kimga bermang.").font = Font(name="Calibri", bold=True, color="C62828")

    # ---- 2) Foydalanuvchilar ----
    ws = wb.create_sheet("Foydalanuvchilar")
    rows = []
    for i, u in enumerate(sorted(users, key=lambda u: u.get("created_at") or datetime.min, reverse=True), 1):
        te = u.get("trial_ends_at")
        rows.append([
            i, owner_label(u), u.get("telegram_id"), u.get("login"),
            to_local(u.get("created_at")), to_local(te),
            "Faol" if te and te > now_utc else "Muddati tugagan",
            accs_per_user.get(u["id"], 0), scans_per_user.get(u["id"], 0),
        ])
    write_table(ws, ["№", "Foydalanuvchi", "Telegram ID", "Login", "Ro'yxatdan o'tgan", "Obuna tugaydi",
                     "Holati", "Akkauntlar", "Skanerlar"], rows,
                [6, 26, 16, 22, 21, 21, 17, 12, 12], status_col=7)

    # ---- 3) Hero akkauntlar ----
    ws = wb.create_sheet("Hero akkauntlar")
    ordered = sorted(accs, key=lambda a: (owner_label(by_id.get(a["user_id"]), a["user_id"]).lower(),
                                          (a.get("email") or "").lower()))
    rows = []
    for i, a in enumerate(ordered, 1):
        n = email_count[(a.get("email") or "").strip().lower()]
        rows.append([
            i, owner_label(by_id.get(a["user_id"]), a["user_id"]), a.get("email"), a.get("hero_password"),
            "Aktiv" if is_active_token(a.get("bearer_token")) else "Nofaol",
            f"Takror ×{n}" if n > 1 else "—",
        ])
    write_table(ws, ["№", "Egasi", "Hero email", "Hero parol", "Holati", "Takror"], rows,
                [6, 26, 34, 24, 12, 14], status_col=5, warn_col=6)

    # ---- 4) Skaner tarixi ----
    ws = wb.create_sheet("Skaner tarixi")
    rows = []
    for i, s in enumerate(sorted(scans, key=lambda s: s.get("scanned_at") or datetime.min, reverse=True), 1):
        ok, total = s.get("success_count") or 0, s.get("total_count") or 0
        rows.append([
            i, owner_label(by_id.get(s["user_id"]), s["user_id"]),
            f"{ok}/{total}" if total else f"{ok} ta",
            int(round(100 * ok / total)) if total else "—",
            round(float(s.get("duration") or 0), 2),
            to_local(s.get("scanned_at")),
        ])
    write_table(ws, ["№", "Foydalanuvchi", "Natija", "Foiz, %", "Davomiyligi, sek", "Sana va vaqt"], rows,
                [6, 26, 12, 10, 18, 22])

    # ---- 5) Arxiv ----
    ws = wb.create_sheet("Arxiv")
    rows = []
    for i, a in enumerate(sorted(arch, key=lambda a: a.get("deleted_at") or datetime.min, reverse=True), 1):
        rows.append([i, owner_label(by_id.get(a["user_id"]), a["user_id"]), a.get("email"),
                     a.get("hero_password"), to_local(a.get("deleted_at"))])
    write_table(ws, ["№", "Egasi", "Hero email", "Hero parol", "O'chirilgan vaqti"], rows,
                [6, 26, 34, 24, 22])

    # ---- 6) Ulashish kodlari ----
    ws = wb.create_sheet("Ulashish kodlari")
    rows = []
    for i, sc in enumerate(sorted(shares, key=lambda x: x.get("created_at") or datetime.min, reverse=True), 1):
        rows.append([
            i, sc.get("code"), owner_label(by_id.get(sc["from_user_id"]), sc["from_user_id"]),
            to_local(sc.get("expires_at")), "Ha" if sc.get("used") else "Yo'q",
            owner_label(by_id.get(sc["used_by"]), sc["used_by"]) if sc.get("used_by") else "—",
            to_local(sc.get("created_at")),
        ])
    write_table(ws, ["№", "Kod", "Yaratgan", "Amal qilish muddati", "Ishlatilgan", "Kim ishlatgan", "Yaratilgan"], rows,
                [6, 14, 26, 22, 13, 26, 22])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _json_default(o):
    if isinstance(o, (datetime, date)):
        return o.isoformat()
    return str(o)


def build_json(data, now_utc, now_local, stats):
    def dump(include_details):
        scan_rows = data["scan_logs"] if include_details else [
            {k: v for k, v in s.items() if k != "details"} for s in data["scan_logs"]]
        payload = {
            "meta": {
                "format": "heroscanner-backup",
                "version": 1,
                "created_at_utc": now_utc.isoformat(),
                "created_at_tashkent": now_local.isoformat(),
                "counts": stats,
                "notes": [
                    "Barcha vaqtlar UTC (bazada saqlanganidek). Toshkent vaqti = UTC + 5 soat.",
                    "app_users.password — bcrypt hash (shu hash bilan foydalanuvchi kirishi tiklanadi).",
                    "hero_accounts.hero_password va archived_accounts.hero_password — shifrdan OCHILGAN (oddiy matn). "
                    "Tiklashda ENCRYPTION_KEY bilan qayta shifrlang.",
                    "bearer_token saqlanmagan — tiklangandan keyin tokenlar avtomatik/qo'lda yangilanadi.",
                    "Fayl maxfiy: parollar bor, hech kimga bermang.",
                ],
            },
            "app_users": data["app_users"],
            "hero_accounts": [{k: v for k, v in a.items() if k != "bearer_token"} for a in data["hero_accounts"]],
            "scan_logs": scan_rows,
            "archived_accounts": data["archived_accounts"],
            "share_codes": data["share_codes"],
        }
        return json.dumps(payload, ensure_ascii=False, indent=2, default=_json_default).encode("utf-8")

    raw = dump(True)
    if len(raw) > MAX_JSON_BYTES:   # juda katta bo'lib ketsa, skaner tafsilotlarini tashlab yuboramiz
        raw = dump(False)
    return raw


def build_files(data):
    """-> (xlsx_bytes, json_bytes, stats, now_local)"""
    now_utc = utc_now()
    now_local = now_utc + LOCAL_OFFSET
    stats = summarize(data, now_utc)
    return build_xlsx(data, now_utc, now_local, stats), build_json(data, now_utc, now_local, stats), stats, now_local
